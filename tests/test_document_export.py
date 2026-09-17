import hashlib
import tempfile
import unittest
from pathlib import Path

from docx import Document

from src.brief_rendering import DecisionBriefRenderer, EvidencePackRenderer, safe_filename
from src.decision_briefs import DecisionBrief, EvidencePack
from src.persistence import SQLitePersistenceStore


class TestDocumentExport(unittest.TestCase):
    def setUp(self):
        self.pack = EvidencePack(
            "Evidence export", "scenario", "scenario-1", [{"evidence_id": "evidence-1"}],
            evidence_gaps=[{"gap_id": "gap-1"}], conflicts=[{"conflict_id": "conflict-1"}],
            programme_references=["programme-1"], mechanism_references=["mechanism-1"],
            outcome_metrics=[{"metric_id": "metric-1"}], provenance="source-register",
            known_limitations=["Limited evidence"], generated_at="2026-01-01T00:00:00Z",
        ).finalized()
        self.brief = DecisionBrief(
            "Brief export", "scenario", "scenario-1", "Human review only", "Strategic question",
            self.pack.pack_id, scenarios_considered=[{"scenario_id": "scenario-1"}],
            alternatives_considered=["Alternative A"], reviews=[{"review_id": "review-1"}],
            dissent_or_reservations=["Reservation"], unresolved_questions=["Question"],
            active_decision={"decision_id": "decision-1"}, authorization_state=[{"authorization_id": "auth-1"}],
            key_uncertainties=["Uncertain"], provenance="brief-register", generated_at="2026-01-02T00:00:00Z",
        ).finalized()

    def test_docx_is_parseable_traceable_and_overwrite_safe(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / safe_filename("decision_brief", self.brief.brief_id, "docx")
            record = DecisionBriefRenderer().write_docx(self.brief, path)
            self.assertTrue(path.is_file())
            self.assertEqual(record.checksum, hashlib.sha256(path.read_bytes()).hexdigest())
            self.assertEqual(record.entity_id, self.brief.brief_id)
            self.assertEqual(record.version, 1)
            document = Document(path)
            text = "\n".join(p.text for p in document.paragraphs)
            table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            for expected in (self.brief.brief_id, self.pack.pack_id, "Dissent or reservations",
                             "Reservation", "Responsible-AI notice", "not a recommendation"):
                self.assertIn(expected, text + table_text)
            with self.assertRaises(FileExistsError):
                DecisionBriefRenderer().write_docx(self.brief, path)

    def test_pdf_is_local_traceable_and_overwrite_safe(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / safe_filename("decision_brief", self.brief.brief_id, "pdf")
            record = DecisionBriefRenderer().write_pdf(self.brief, path)
            data = path.read_bytes()
            self.assertTrue(data.startswith(b"%PDF-"))
            self.assertGreater(len(data), 500)
            self.assertEqual(record.checksum, hashlib.sha256(data).hexdigest())
            self.assertEqual(record.source_snapshot_id, self.brief.brief_id)
            for expected in (self.brief.brief_id, self.pack.pack_id, "Dissent or reservations",
                             "Reservation", "not a recommendation"):
                self.assertIn(expected.encode("utf-8"), data)
            with self.assertRaises(FileExistsError):
                DecisionBriefRenderer().write_pdf(self.brief, path)

    def test_pack_uses_same_normalized_sections_and_export_audits_are_factual(self):
        brief_sections = [heading for heading, _ in DecisionBriefRenderer().normalized_document(self.brief)["sections"]]
        pack_sections = [heading for heading, _ in EvidencePackRenderer().normalized_document(self.pack)["sections"]]
        self.assertIn("Evidence gaps", pack_sections)
        self.assertIn("Dissent or reservations", brief_sections)
        self.assertNotIn("winner", " ".join(brief_sections).lower())
        with tempfile.TemporaryDirectory() as directory:
            db = Path(directory) / "exports.sqlite3"
            with SQLitePersistenceStore(db) as store:
                EvidencePackRenderer().write_docx(self.pack, Path(directory) / "pack.docx", audit_events=store.audit_events)
                DecisionBriefRenderer().write_pdf(self.brief, Path(directory) / "brief.pdf", audit_events=store.audit_events)
                events = store.audit_events.list_events()
                docx = next(event for event in events if event.event_type == "evidence_pack_docx_exported")
                pdf = next(event for event in events if event.event_type == "decision_brief_pdf_exported")
                self.assertEqual(docx.entity_id, self.pack.pack_id)
                self.assertEqual(pdf.metadata["version"], 1)
                self.assertIn("checksum", pdf.metadata)

