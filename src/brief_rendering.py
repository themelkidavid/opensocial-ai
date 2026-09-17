"""Deterministic, local renderers for traceable evidence packs and decision briefs."""

from dataclasses import dataclass, asdict
from datetime import datetime, timezone
from hashlib import sha256
from html import escape
import json
from pathlib import Path
import re
from typing import Any, Iterable, Optional, Union

try:
    from src.decision_briefs import DecisionBrief, EvidencePack
except ModuleNotFoundError:
    from decision_briefs import DecisionBrief, EvidencePack


_NOTICE = (
    "This rendered summary is for human review. It is not a recommendation, "
    "prediction, authorization, or evidence of effectiveness."
)


@dataclass(frozen=True)
class ExportRecord:
    """Factual metadata for one locally rendered export."""

    entity_type: str
    entity_id: str
    format: str
    checksum: str
    rendered_at: str
    path: Optional[str] = None
    subject_type: Optional[str] = None
    subject_id: Optional[str] = None
    source_snapshot_id: Optional[str] = None
    version: Optional[int] = None

    def to_dict(self) -> dict:
        return asdict(self)


def safe_filename(entity_type: str, entity_id: str, extension: str) -> str:
    """Return a portable filename without allowing caller-controlled paths."""

    def clean(value: str) -> str:
        value = re.sub(r"[^A-Za-z0-9._-]+", "-", str(value)).strip(".-")
        return value or "unknown"

    return f"{clean(entity_type)}-{clean(entity_id)}.{clean(extension)}"


def _value(value: Any) -> str:
    if value is None or value == "":
        return "Not available"
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
    return str(value)


def _sections(items: Iterable[tuple[str, Any]]) -> list[tuple[str, str]]:
    return [(title, _value(value)) for title, value in items if value not in (None, "", [], {})]


class _Renderer:
    entity_type = ""

    def _audit(self, audit_events: Any, event_type: str, entity_id: str, subject_type: str,
               subject_id: str, description: str, metadata: dict) -> None:
        if audit_events is not None:
            audit_events.record(
                event_type, self.entity_type, entity_id, description,
                {"subject_type": subject_type, "subject_id": subject_id, **metadata},
            )

    def _markdown(self, title: str, metadata: list[tuple[str, Any]],
                  sections: list[tuple[str, str]], notice: bool) -> str:
        lines = [f"# {title}", "", "## Traceability"]
        lines.extend(f"- **{label}:** {_value(value)}" for label, value in metadata)
        if notice:
            lines.extend(["", "## Responsible-AI notice", _NOTICE])
        for heading, value in sections:
            lines.extend(["", f"## {heading}", value])
        return "\n".join(lines) + "\n"

    def _html(self, title: str, metadata: list[tuple[str, Any]],
              sections: list[tuple[str, str]], notice: bool) -> str:
        body = ["<!doctype html>", '<html lang="en"><head><meta charset="utf-8">',
                f"<title>{escape(title)}</title></head><body>", f"<h1>{escape(title)}</h1>",
                "<h2>Traceability</h2><dl>"]
        for label, value in metadata:
            body.append(f"<dt>{escape(label)}</dt><dd>{escape(_value(value))}</dd>")
        body.append("</dl>")
        if notice:
            body.append(f"<h2>Responsible-AI notice</h2><p>{escape(_NOTICE)}</p>")
        for heading, value in sections:
            body.append(f"<h2>{escape(heading)}</h2><pre>{escape(value)}</pre>")
        body.append("</body></html>\n")
        return "\n".join(body)

    def write(self, content: str, path: Union[Path, str], overwrite: bool = False) -> ExportRecord:
        destination = Path(path)
        if not destination.parent.is_dir():
            raise ValueError("Export directory does not exist")
        if destination.exists() and not overwrite:
            raise FileExistsError(f"Refusing to overwrite {destination}")
        destination.write_text(content, encoding="utf-8")
        return ExportRecord(
            self.entity_type, "", destination.suffix.lstrip("."),
            sha256(content.encode("utf-8")).hexdigest(),
            datetime.now(timezone.utc).isoformat(), str(destination),
        )

    def normalized_document(self, entity: Any, include_responsible_ai_notice: bool = True) -> dict:
        """Return the one ordered, source-faithful document representation for all formats."""
        entity, metadata, sections = self._data(entity)
        return {"title": entity.title, "entity_type": self.entity_type,
                "entity_id": getattr(entity, "brief_id", None) or getattr(entity, "pack_id", None),
                "subject_type": entity.subject_type, "subject_id": entity.subject_id,
                "version": getattr(entity, "brief_version", None), "metadata": metadata,
                "sections": sections,
                "responsible_ai_notice": _NOTICE if include_responsible_ai_notice else None}

    def _export_record(self, document: dict, path: Path) -> ExportRecord:
        return ExportRecord(document["entity_type"], document["entity_id"], path.suffix.lstrip("."),
                            sha256(path.read_bytes()).hexdigest(), datetime.now(timezone.utc).isoformat(), str(path),
                            document["subject_type"], document["subject_id"], document["entity_id"],
                            document["version"])


class DocumentExporter:
    """Shared safe output boundary; format subclasses receive normalized sections only."""

    extension = ""

    def _destination(self, path: Union[Path, str], overwrite: bool) -> Path:
        destination = Path(path)
        if not destination.parent.is_dir():
            raise ValueError("Export directory does not exist")
        if destination.exists() and not overwrite:
            raise FileExistsError(f"Refusing to overwrite {destination}")
        if destination.suffix.lower() != f".{self.extension}":
            raise ValueError(f"Export path must end in .{self.extension}")
        return destination


class DocxExporter(DocumentExporter):
    """Local DOCX formatter for normalized renderer output."""

    extension = "docx"

    def write(self, renderer: _Renderer, entity: Any, path: Union[Path, str], overwrite: bool = False,
              include_responsible_ai_notice: bool = True, audit_events: Any = None) -> ExportRecord:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt
        destination, data = self._destination(path, overwrite), renderer.normalized_document(entity, include_responsible_ai_notice)
        document = Document(); section = document.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.7)
        document.add_heading(data["title"], 0)
        table = document.add_table(rows=0, cols=2); table.style = "Table Grid"
        for label, value in data["metadata"]:
            row = table.add_row().cells; row[0].text, row[1].text = label, _value(value)
        if data["responsible_ai_notice"]:
            document.add_heading("Responsible-AI notice", 1); document.add_paragraph(data["responsible_ai_notice"])
        for heading, value in data["sections"]:
            document.add_heading(heading, 1)
            try: parsed = json.loads(value)
            except (TypeError, ValueError): parsed = None
            if isinstance(parsed, list):
                for item in parsed: document.add_paragraph(_value(item), style="List Bullet")
            else: document.add_paragraph(value)
        footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.text = f"{data['entity_type']} ID: {data['entity_id']}" + (f" | Version: {data['version']}" if data["version"] is not None else "")
        for paragraph in document.paragraphs:
            for run in paragraph.runs: run.font.size = Pt(10)
        document.core_properties.title = data["title"]
        document.core_properties.subject = f"{data['subject_type']}: {data['subject_id']}"
        document.core_properties.comments = "Traceable local export; checksum is not a digital signature."
        document.save(destination)
        record = renderer._export_record(data, destination)
        renderer._audit(audit_events, f"{data['entity_type']}_docx_exported", data["entity_id"], data["subject_type"], data["subject_id"], "Document exported as DOCX.", {"format": "docx", "filename": destination.name, "checksum": record.checksum, "version": data["version"]})
        return record


class PdfExporter(DocumentExporter):
    """Browser-free local PDF formatter for normalized renderer output."""

    extension = "pdf"

    def write(self, renderer: _Renderer, entity: Any, path: Union[Path, str], overwrite: bool = False,
              include_responsible_ai_notice: bool = True, audit_events: Any = None) -> ExportRecord:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Preformatted
        destination, data = self._destination(path, overwrite), renderer.normalized_document(entity, include_responsible_ai_notice)
        styles = getSampleStyleSheet(); styles.add(ParagraphStyle(name="DocumentTitle", parent=styles["Title"], alignment=TA_CENTER)); body = styles["BodyText"]
        story = [Paragraph(escape(data["title"]), styles["DocumentTitle"]), Spacer(1, 12)]
        rows = [[Paragraph(f"<b>{escape(label)}</b>", body), Paragraph(escape(_value(value)), body)] for label, value in data["metadata"]]
        table = Table(rows, colWidths=[1.65 * inch, 5.35 * inch]); table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (0, 0), (0, -1), colors.whitesmoke), ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6)]))
        story.extend([table, Spacer(1, 12)])
        if data["responsible_ai_notice"]: story.extend([Paragraph("Responsible-AI notice", styles["Heading2"]), Paragraph(escape(data["responsible_ai_notice"]), body), Spacer(1, 8)])
        for heading, value in data["sections"]: story.extend([Paragraph(escape(heading), styles["Heading2"]), Preformatted(value, styles["Code"]), Spacer(1, 8)])
        def footer(canvas, doc):
            canvas.saveState(); canvas.setFont("Helvetica", 8)
            text = f"{data['entity_type']} ID: {data['entity_id']}" + (f" | Version: {data['version']}" if data["version"] is not None else "")
            canvas.drawCentredString(letter[0] / 2, 0.45 * inch, text + f" | Page {doc.page}"); canvas.restoreState()
        pdf = SimpleDocTemplate(str(destination), pagesize=letter, leftMargin=0.7 * inch, rightMargin=0.7 * inch, topMargin=0.7 * inch, bottomMargin=0.7 * inch, pageCompression=0)
        pdf.title = data["title"]; pdf.build(story, onFirstPage=footer, onLaterPages=footer)
        record = renderer._export_record(data, destination)
        renderer._audit(audit_events, f"{data['entity_type']}_pdf_exported", data["entity_id"], data["subject_type"], data["subject_id"], "Document exported as PDF.", {"format": "pdf", "filename": destination.name, "checksum": record.checksum, "version": data["version"]})
        return record


class EvidencePackRenderer(_Renderer):
    entity_type = "evidence_pack"

    def _data(self, pack: EvidencePack) -> tuple[EvidencePack, list, list]:
        if not isinstance(pack, EvidencePack):
            raise TypeError("pack must be an EvidencePack")
        pack = pack.finalized()
        metadata = [("Evidence pack ID", pack.pack_id), ("Subject type", pack.subject_type),
                    ("Subject ID", pack.subject_id), ("Generated at", pack.generated_at),
                    ("Provenance", pack.provenance)]
        sections = _sections([
            ("Evidence records", pack.evidence_records), ("Evidence quality summary", pack.evidence_quality_summary),
            ("Evidence gaps", pack.evidence_gaps), ("Conflicts", pack.conflicts),
            ("Patterns", pack.patterns), ("Insights", pack.insights),
            ("Programme references", pack.programme_references), ("Mechanism references", pack.mechanism_references),
            ("Outcome metrics", pack.outcome_metrics), ("Observations", pack.observations),
            ("Source hypothesis IDs", pack.source_hypothesis_ids),
            ("Known limitations", pack.known_limitations),
        ])
        return pack, metadata, sections

    def render_markdown(self, pack: EvidencePack, include_responsible_ai_notice: bool = True,
                        audit_events: Any = None) -> str:
        pack, metadata, sections = self._data(pack)
        result = self._markdown(pack.title, metadata, sections, include_responsible_ai_notice)
        self._audit(audit_events, "evidence_pack_rendered", pack.pack_id, pack.subject_type, pack.subject_id,
                    "Evidence pack rendered", {"format": "markdown"})
        return result

    def render_html(self, pack: EvidencePack, include_responsible_ai_notice: bool = True,
                    audit_events: Any = None) -> str:
        pack, metadata, sections = self._data(pack)
        result = self._html(pack.title, metadata, sections, include_responsible_ai_notice)
        self._audit(audit_events, "evidence_pack_rendered", pack.pack_id, pack.subject_type, pack.subject_id,
                    "Evidence pack rendered", {"format": "html"})
        return result

    def write_markdown(self, pack: EvidencePack, path: Union[Path, str], overwrite: bool = False,
                       audit_events: Any = None) -> ExportRecord:
        pack = pack.finalized()
        record = self.write(self.render_markdown(pack), path, overwrite)
        record = ExportRecord(record.entity_type, pack.pack_id, record.format, record.checksum,
                              record.rendered_at, record.path, pack.subject_type, pack.subject_id,
                              pack.pack_id)
        self._audit(audit_events, "evidence_pack_exported", pack.pack_id, pack.subject_type, pack.subject_id,
                    "Evidence pack exported", {"format": "markdown"})
        return record

    def write_html(self, pack: EvidencePack, path: Union[Path, str], overwrite: bool = False,
                   audit_events: Any = None) -> ExportRecord:
        pack = pack.finalized()
        record = self.write(self.render_html(pack), path, overwrite)
        record = ExportRecord(record.entity_type, pack.pack_id, record.format, record.checksum,
                              record.rendered_at, record.path, pack.subject_type, pack.subject_id,
                              pack.pack_id)
        self._audit(audit_events, "evidence_pack_exported", pack.pack_id, pack.subject_type, pack.subject_id,
                    "Evidence pack exported", {"format": "html"})
        return record

    def write_docx(self, pack: EvidencePack, path: Union[Path, str], overwrite: bool = False,
                   include_responsible_ai_notice: bool = True, audit_events: Any = None) -> ExportRecord:
        return DocxExporter().write(self, pack, path, overwrite, include_responsible_ai_notice, audit_events)

    def write_pdf(self, pack: EvidencePack, path: Union[Path, str], overwrite: bool = False,
                  include_responsible_ai_notice: bool = True, audit_events: Any = None) -> ExportRecord:
        return PdfExporter().write(self, pack, path, overwrite, include_responsible_ai_notice, audit_events)


class DecisionBriefRenderer(_Renderer):
    entity_type = "decision_brief"

    def _data(self, brief: DecisionBrief) -> tuple[DecisionBrief, list, list]:
        if not isinstance(brief, DecisionBrief):
            raise TypeError("brief must be a DecisionBrief")
        brief = brief.finalized()
        metadata = [("Decision brief ID", brief.brief_id), ("Version", brief.brief_version),
                    ("Subject type", brief.subject_type), ("Subject ID", brief.subject_id),
                    ("Evidence pack ID", brief.evidence_pack_id), ("Generated at", brief.generated_at),
                    ("Provenance", brief.provenance)]
        sections = _sections([
            ("Purpose", brief.purpose), ("Problem summary", brief.problem_summary),
            ("Scenarios considered", brief.scenarios_considered), ("Alternatives considered", brief.alternatives_considered),
            ("Human reviews", brief.reviews), ("Concerns", brief.concerns),
            ("Dissent or reservations", brief.dissent_or_reservations),
            ("Unresolved questions", brief.unresolved_questions), ("Current human decision", brief.active_decision),
            ("Authorization state", brief.authorization_state), ("Expected learning", brief.expected_learning),
            ("Key uncertainties", brief.key_uncertainties),
            ("Source hypothesis IDs", brief.source_hypothesis_ids),
        ])
        return brief, metadata, sections

    def render_markdown(self, brief: DecisionBrief, include_responsible_ai_notice: bool = True,
                        audit_events: Any = None) -> str:
        brief, metadata, sections = self._data(brief)
        result = self._markdown(brief.title, metadata, sections, include_responsible_ai_notice)
        self._audit(audit_events, "decision_brief_rendered", brief.brief_id, brief.subject_type, brief.subject_id,
                    "Decision brief rendered", {"format": "markdown", "version": brief.brief_version})
        return result

    def render_html(self, brief: DecisionBrief, include_responsible_ai_notice: bool = True,
                    audit_events: Any = None) -> str:
        brief, metadata, sections = self._data(brief)
        result = self._html(brief.title, metadata, sections, include_responsible_ai_notice)
        self._audit(audit_events, "decision_brief_rendered", brief.brief_id, brief.subject_type, brief.subject_id,
                    "Decision brief rendered", {"format": "html", "version": brief.brief_version})
        return result

    def write_markdown(self, brief: DecisionBrief, path: Union[Path, str], overwrite: bool = False,
                       audit_events: Any = None) -> ExportRecord:
        brief = brief.finalized()
        record = self.write(self.render_markdown(brief), path, overwrite)
        record = ExportRecord(record.entity_type, brief.brief_id, record.format, record.checksum,
                              record.rendered_at, record.path, brief.subject_type, brief.subject_id,
                              brief.brief_id, brief.brief_version)
        self._audit(audit_events, "decision_brief_exported", brief.brief_id, brief.subject_type, brief.subject_id,
                    "Decision brief exported", {"format": "markdown", "version": brief.brief_version})
        return record

    def write_html(self, brief: DecisionBrief, path: Union[Path, str], overwrite: bool = False,
                   audit_events: Any = None) -> ExportRecord:
        brief = brief.finalized()
        record = self.write(self.render_html(brief), path, overwrite)
        record = ExportRecord(record.entity_type, brief.brief_id, record.format, record.checksum,
                              record.rendered_at, record.path, brief.subject_type, brief.subject_id,
                              brief.brief_id, brief.brief_version)
        self._audit(audit_events, "decision_brief_exported", brief.brief_id, brief.subject_type, brief.subject_id,
                    "Decision brief exported", {"format": "html", "version": brief.brief_version})
        return record

    def write_docx(self, brief: DecisionBrief, path: Union[Path, str], overwrite: bool = False,
                   include_responsible_ai_notice: bool = True, audit_events: Any = None) -> ExportRecord:
        return DocxExporter().write(self, brief, path, overwrite, include_responsible_ai_notice, audit_events)

    def write_pdf(self, brief: DecisionBrief, path: Union[Path, str], overwrite: bool = False,
                  include_responsible_ai_notice: bool = True, audit_events: Any = None) -> ExportRecord:
        return PdfExporter().write(self, brief, path, overwrite, include_responsible_ai_notice, audit_events)
